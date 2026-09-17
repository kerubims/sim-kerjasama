#!/usr/bin/env python3
import sys
import os
import json
import zipfile
import base64
import re
from xml.etree import ElementTree as ET
import docx
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

NAMESPACES = {
    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
    'pic': 'http://schemas.openxmlformats.org/drawingml/2006/picture'
}

for prefix, uri in NAMESPACES.items():
    ET.register_namespace(prefix, uri)

def parse_docx_to_ast(docx_path):
    if not zipfile.is_zipfile(docx_path):
        raise ValueError("File is not a valid docx zip archive.")

    with zipfile.ZipFile(docx_path, 'r') as z:
        doc_xml = z.read('word/document.xml')
        rels_xml = z.read('word/_rels/document.xml.rels') if 'word/_rels/document.xml.rels' in z.namelist() else None

    # Parse relationships
    rels = {}
    if rels_xml:
        rels_tree = ET.fromstring(rels_xml)
        for rel in rels_tree.findall('{http://schemas.openxmlformats.org/package/2006/relationships}Relationship'):
            rels[rel.attrib.get('Id')] = rel.attrib.get('Target')

    # Read media images as base64
    media = {}
    with zipfile.ZipFile(docx_path, 'r') as z:
        for name in z.namelist():
            if name.startswith('word/media/'):
                img_bytes = z.read(name)
                ext = name.split('.')[-1].lower()
                mime = 'image/png' if ext == 'png' else ('image/jpeg' if ext in ['jpg', 'jpeg'] else 'image/gif')
                b64 = base64.b64encode(img_bytes).decode('utf-8')
                media[name.replace('word/', '')] = f"data:{mime};base64,{b64}"

    tree = ET.fromstring(doc_xml)
    body = tree.find('w:body', NAMESPACES)

    ast = {
        "version": "1.0",
        "sections": [],
        "body": []
    }

    def parse_run(r_elem):
        run_data = {
            "type": "run",
            "text": "",
            "bold": False,
            "italic": False,
            "underline": False,
            "fontFamily": None,
            "fontSize": None,
            "color": None,
            "highlight": None,
            "image": None
        }

        rPr = r_elem.find('w:rPr', NAMESPACES)
        if rPr is not None:
            if rPr.find('w:b', NAMESPACES) is not None:
                run_data["bold"] = True
            if rPr.find('w:i', NAMESPACES) is not None:
                run_data["italic"] = True
            if rPr.find('w:u', NAMESPACES) is not None:
                run_data["underline"] = True
            
            rFonts = rPr.find('w:rFonts', NAMESPACES)
            if rFonts is not None:
                run_data["fontFamily"] = rFonts.attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii') or rFonts.attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hAnsi')
            
            sz = rPr.find('w:sz', NAMESPACES)
            if sz is not None:
                val = sz.attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
                if val:
                    run_data["fontSize"] = float(val) / 2.0  # half-points to pt

            color = rPr.find('w:color', NAMESPACES)
            if color is not None:
                val = color.attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
                if val and val != 'auto':
                    run_data["color"] = f"#{val}"

            highlight = rPr.find('w:highlight', NAMESPACES)
            if highlight is not None:
                run_data["highlight"] = highlight.attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')

        # Text
        t_elem = r_elem.find('w:t', NAMESPACES)
        if t_elem is not None and t_elem.text:
            run_data["text"] = t_elem.text

        # Image (drawing)
        drawing = r_elem.find('w:drawing', NAMESPACES)
        if drawing is not None:
            blip = drawing.find('.//a:blip', NAMESPACES)
            if blip is not None:
                embed_id = blip.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                target = rels.get(embed_id)
                if target and target in media:
                    run_data["image"] = media[target]

        return run_data

    def parse_paragraph(p_elem):
        p_data = {
            "type": "paragraph",
            "alignment": "left",
            "style": "Normal",
            "runs": []
        }

        pPr = p_elem.find('w:pPr', NAMESPACES)
        if pPr is not None:
            jc = pPr.find('w:jc', NAMESPACES)
            if jc is not None:
                p_data["alignment"] = jc.attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', 'left')

            pStyle = pPr.find('w:pStyle', NAMESPACES)
            if pStyle is not None:
                p_data["style"] = pStyle.attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', 'Normal')

        for child in p_elem:
            if child.tag.endswith('r'):
                p_data["runs"].append(parse_run(child))

        return p_data

    def parse_table(tbl_elem):
        tbl_data = {
            "type": "table",
            "rows": []
        }

        for tr in tbl_elem.findall('w:tr', NAMESPACES):
            row_data = {"cells": []}
            for tc in tr.findall('w:tc', NAMESPACES):
                cell_data = {
                    "shading": None,
                    "paragraphs": []
                }

                tcPr = tc.find('w:tcPr', NAMESPACES)
                if tcPr is not None:
                    shd = tcPr.find('w:shd', NAMESPACES)
                    if shd is not None:
                        fill = shd.attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill')
                        if fill and fill != 'auto':
                            cell_data["shading"] = f"#{fill}"

                for p in tc.findall('w:p', NAMESPACES):
                    cell_data["paragraphs"].append(parse_paragraph(p))

                row_data["cells"].append(cell_data)

            tbl_data["rows"].append(row_data)

        return tbl_data

    if body is not None:
        for elem in body:
            if elem.tag.endswith('p'):
                ast["body"].append(parse_paragraph(elem))
            elif elem.tag.endswith('tbl'):
                ast["body"].append(parse_table(elem))

    return ast

def ast_to_docx(ast_data, output_path):
    doc = Document()

    # Page setup - A4
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    for item in ast_data.get("body", []):
        if item.get("type") == "paragraph":
            p = doc.add_paragraph()
            align = item.get("alignment", "left").lower()
            if align == "center":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif align == "right":
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif align == "justify":
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

            for r_data in item.get("runs", []):
                text = r_data.get("text", "")
                if text:
                    run = p.add_run(text)
                    if r_data.get("bold"):
                        run.bold = True
                    if r_data.get("italic"):
                        run.italic = True
                    if r_data.get("underline"):
                        run.underline = True
                    if r_data.get("fontFamily"):
                        run.font.name = r_data["fontFamily"]
                    if r_data.get("fontSize"):
                        run.font.size = Pt(r_data["fontSize"])
                    if r_data.get("color"):
                        c = r_data["color"].lstrip('#')
                        if len(c) == 6:
                            run.font.color.rgb = RGBColor(int(c[0:2], 16), int(c[2:4], 16), int(c[4:6], 16))

        elif item.get("type") == "table":
            rows = item.get("rows", [])
            if not rows:
                continue
            num_rows = len(rows)
            num_cols = max([len(r.get("cells", [])) for r in rows]) if rows else 0
            
            table = doc.add_table(rows=num_rows, cols=num_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER

            for r_idx, r_data in enumerate(rows):
                for c_idx, c_data in enumerate(r_data.get("cells", [])):
                    if c_idx >= num_cols:
                        break
                    cell = table.cell(r_idx, c_idx)
                    shading = c_data.get("shading")
                    if shading:
                        hex_c = shading.lstrip('#')
                        shd_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_c}"/>')
                        cell._tc.get_or_add_tcPr().append(shd_elm)

                    # Paragraphs in cell
                    cell_paras = c_data.get("paragraphs", [])
                    if cell_paras:
                        p = cell.paragraphs[0]
                        for idx, p_data in enumerate(cell_paras):
                            if idx > 0:
                                p = cell.add_paragraph()
                            align = p_data.get("alignment", "left").lower()
                            if align == "center":
                                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            elif align == "right":
                                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                            elif align == "justify":
                                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                            else:
                                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

                            for r_data in p_data.get("runs", []):
                                text = r_data.get("text", "")
                                if text:
                                    run = p.add_run(text)
                                    if r_data.get("bold"):
                                        run.bold = True
                                    if r_data.get("italic"):
                                        run.italic = True
                                    if r_data.get("underline"):
                                        run.underline = True
                                    if r_data.get("fontFamily"):
                                        run.font.name = r_data["fontFamily"]
                                    if r_data.get("fontSize"):
                                        run.font.size = Pt(r_data["fontSize"])

    doc.save(output_path)

if __name__ == '__main__':
    if len(sys.argv) < 3:
        print("Usage: DocxAST.py <command> <input_file> [output_file]")
        print("Commands: parse | build")
        sys.exit(1)

    cmd = sys.argv[1]
    input_file = sys.argv[2]

    if cmd == "parse":
        ast = parse_docx_to_ast(input_file)
        print(json.dumps(ast, indent=2))
    elif cmd == "build":
        if len(sys.argv) < 4:
            print("Output path required for build command.")
            sys.exit(1)
        output_file = sys.argv[3]
        with open(input_file, 'r', encoding='utf-8') as f:
            ast_data = json.load(f)
        ast_to_docx(ast_data, output_file)
        print(f"Built DOCX successfully to {output_file}")
