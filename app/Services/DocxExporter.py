#!/usr/bin/env python3
import sys
import os
import re
import html
from bs4 import BeautifulSoup
import docx
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def parse_css_style(style_str):
    styles = {}
    if not style_str:
        return styles
    for item in style_str.split(';'):
        if ':' in item:
            k, v = item.split(':', 1)
            styles[k.strip().lower()] = v.strip()
    return styles

def parse_color(color_str):
    if not color_str:
        return None
    color_str = color_str.strip().lower()
    if color_str.startswith('#'):
        hex_val = color_str.lstrip('#')
        if len(hex_val) == 3:
            hex_val = ''.join([c*2 for c in hex_val])
        if len(hex_val) == 6:
            try:
                return RGBColor(int(hex_val[0:2], 16), int(hex_val[2:4], 16), int(hex_val[4:6], 16))
            except ValueError:
                return None
    elif color_str.startswith('rgb'):
        nums = re.findall(r'\d+', color_str)
        if len(nums) >= 3:
            return RGBColor(int(nums[0]), int(nums[1]), int(nums[2]))
    return None

def set_cell_background(cell, hex_color):
    if not hex_color:
        return
    hex_color = hex_color.lstrip('#')
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def html_to_docx(html_content, output_path):
    doc = Document()
    
    # Page setup - A4
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    soup = BeautifulSoup(html_content, 'html.parser')

    def process_node(node, parent_p=None, current_styles=None):
        if current_styles is None:
            current_styles = {}

        if node.name is None:  # Text node
            text_val = str(node).strip()
            if text_val:
                if parent_p is None:
                    parent_p = doc.add_paragraph()
                run = parent_p.add_run(str(node))
                
                font_family = current_styles.get('font-family')
                if font_family:
                    clean_font = font_family.split(',')[0].strip(' "\'')
                    run.font.name = clean_font

                font_size = current_styles.get('font-size')
                if font_size:
                    if font_size.endswith('px'):
                        px = float(font_size.replace('px', ''))
                        run.font.size = Pt(px * 0.75)
                    elif font_size.endswith('pt'):
                        pt = float(font_size.replace('pt', ''))
                        run.font.size = Pt(pt)

                font_color = current_styles.get('color')
                if font_color:
                    rgb = parse_color(font_color)
                    if rgb:
                        run.font.color.rgb = rgb

                if current_styles.get('font-weight') == 'bold':
                    run.bold = True
                if current_styles.get('font-style') == 'italic':
                    run.italic = True
                if current_styles.get('text-decoration') == 'underline':
                    run.underline = True

        elif node.name in ['p', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6']:
            p = doc.add_paragraph()
            p_style = parse_css_style(node.get('style', ''))
            
            align = p_style.get('text-align', '').lower()
            if align == 'center':
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif align == 'right':
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif align == 'justify':
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

            for child in node.children:
                process_node(child, p, p_style.copy())

        elif node.name == 'table':
            table_style = parse_css_style(node.get('style', ''))
            rows = node.find_all('tr', recursive=True)
            if not rows:
                return
            
            num_rows = len(rows)
            max_cols = max([len(r.find_all(['td', 'th'], recursive=True)) for r in rows]) if rows else 0
            if num_cols := max_cols:
                table = doc.add_table(rows=num_rows, cols=num_cols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = False

                for r_idx, tr in enumerate(rows):
                    cells = tr.find_all(['td', 'th'], recursive=False)
                    for c_idx, td in enumerate(cells):
                        if c_idx >= num_cols:
                            break
                        cell = table.cell(r_idx, c_idx)
                        cell_style = parse_css_style(td.get('style', ''))
                        
                        bg_color = cell_style.get('background-color') or cell_style.get('background')
                        if bg_color:
                            set_cell_background(cell, bg_color)
                        
                        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)

                        cell_p = cell.paragraphs[0]
                        for child in td.children:
                            process_node(child, cell_p, cell_style.copy())

        elif node.name in ['span', 'strong', 'b', 'em', 'i', 'u', 'font', 'a']:
            merged_styles = current_styles.copy()
            merged_styles.update(parse_css_style(node.get('style', '')))
            
            if node.name in ['strong', 'b']:
                merged_styles['font-weight'] = 'bold'
            if node.name in ['em', 'i']:
                merged_styles['font-style'] = 'italic'
            if node.name == 'u':
                merged_styles['text-decoration'] = 'underline'

            for child in node.children:
                process_node(child, parent_p, merged_styles)

        else:
            # Container tags like div, section, article, main, header, etc.
            container_styles = current_styles.copy()
            if hasattr(node, 'get'):
                container_styles.update(parse_css_style(node.get('style', '')))

            for child in getattr(node, 'children', []):
                process_node(child, parent_p, container_styles)

    for child in soup.children:
        process_node(child)

    doc.save(output_path)

if __name__ == '__main__':
    if len(sys.argv) < 3:
        print("Usage: DocxExporter.py <html_file> <output_docx>")
        sys.exit(1)
    
    html_file = sys.argv[1]
    output_docx = sys.argv[2]
    
    with open(html_file, 'r', encoding='utf-8') as f:
        html_str = f.read()
    
    html_to_docx(html_str, output_docx)
    print("Exported DOCX successfully to", output_docx)
